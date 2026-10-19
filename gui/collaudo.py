"""Collaudo automatico dell'interfaccia vera, utilizzabile anche sull'eseguibile installato.

    PDFListinoExtractor.exe --collaudo <pdf> <file.xlsx>

Apre la finestra principale, carica il PDF, estrae col processo separato come farebbe l'utente,
esporta in Excel, salva uno screenshot e un report JSON accanto al file Excel, poi chiude.
Codice di uscita 0 se tutto e' andato a buon fine. Non cambia nulla nell'uso normale.
"""
import itertools
import json
import sys
import time
import traceback
from pathlib import Path

OPZIONE = "--collaudo"
# PDFListinoExtractor.exe --collaudo-wizard <pdf> <regole.json>: il wizard dall'analisi al salvataggio,
# scrivendo nel file indicato e non nelle regole personali dell'utente
OPZIONE_WIZARD = "--collaudo-wizard"


def collaudo_richiesto(argv: list[str]) -> bool:
    return OPZIONE in argv or OPZIONE_WIZARD in argv


def esegui_collaudo(argv: list[str]) -> int:
    if OPZIONE_WIZARD in argv:
        return esegui_collaudo_wizard(argv)
    i = argv.index(OPZIONE)
    pdf, xlsx = Path(argv[i + 1]), Path(argv[i + 2])
    report_path = xlsx.with_suffix(".collaudo.json")
    report: dict = {"pdf": str(pdf), "xlsx": str(xlsx), "esito": "fallito"}
    try:
        _collauda(pdf, xlsx, report)
    except Exception as e:
        report["errore"] = f"{type(e).__name__}: {e}"
        report["traceback"] = traceback.format_exc()
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["esito"] == "ok" else 1


def _collauda(pdf: Path, xlsx: Path, report: dict) -> None:
    from openpyxl import load_workbook
    from PIL import ImageGrab

    from exporters.excel_exporter import esporta_listino
    from extractors.pdf_reader import PdfNonValido, verifica_pdf
    from gui.main_window import FinestraPrincipale
    from percorsi import cartella_log, impacchettato
    from version import __version__

    report.update(versione=__version__, impacchettato=impacchettato(), eseguibile=sys.executable,
                  log=str(cartella_log() / "app.log"))
    app = FinestraPrincipale()
    app.geometry("1400x860+20+20")
    report["titolo_finestra"] = app.title()
    report["fornitori"] = list(app.fornitori)
    try:
        app.info_pdf = verifica_pdf(pdf)
    except PdfNonValido as e:
        # anche un PDF non valido e' un esito di collaudo: il messaggio deve essere leggibile
        report.update(esito="pdf_non_valido", messaggio=str(e))
        app.destroy()
        return
    app._aggiorna_pulsanti()
    battiti: list[float] = []
    tempi: dict[str, float] = {}

    def battito():
        battiti.append(time.perf_counter())
        if app._lavoro is not None:
            app.after(50, battito)

    def avvia():
        tempi["inizio"] = time.perf_counter()
        app._avvia_estrazione()
        battito()

    def fine():
        if app._lavoro is not None or "inizio" not in tempi:
            app.after(100, fine)
            return
        try:
            tempi["fine"] = time.perf_counter()
            app.update()
            articoli = app.tabella.articoli()
            n_esportate = esporta_listino(articoli, xlsx, app.tabella.colonne)
            ws = load_workbook(xlsx)["Listino"]
            x, y = app.winfo_rootx(), app.winfo_rooty()
            ImageGrab.grab((x, y, x + app.winfo_width(), y + app.winfo_height())).save(
                xlsx.with_suffix(".collaudo.png")
            )
            pause = [round((b - a) * 1000) for a, b in itertools.pairwise(battiti)]
            report.update(
                stato=app.lbl_stato.cget("text"),
                articoli=len(articoli),
                da_verificare=sum(1 for a in articoli if a.tutte_le_anomalie()),
                righe_esportate=n_esportate,
                righe_excel=ws.max_row - 1,
                intestazioni_excel=[c.value for c in ws[1]],
                prima_riga_excel=[c.value for c in ws[2]] if ws.max_row > 1 else None,
                durata_s=round(tempi["fine"] - tempi["inizio"], 2),
                pausa_massima_ms=max(pause, default=0),
            )
            ok = articoli and n_esportate == len(articoli) == ws.max_row - 1
            report["esito"] = "ok" if ok else "fallito"
        except Exception as e:
            report["errore"] = f"{type(e).__name__}: {e}"
            report["traceback"] = traceback.format_exc()
        finally:
            app.destroy()

    app.after(300, avvia)
    app.after(600, fine)
    app.mainloop()


def esegui_collaudo_wizard(argv: list[str]) -> int:
    i = argv.index(OPZIONE_WIZARD)
    pdf, regole = Path(argv[i + 1]), Path(argv[i + 2])
    # pagina facoltativa: su un catalogo la prima pagina e' spesso l'indice, non una tabella
    pagina = int(argv[i + 3]) if len(argv) > i + 3 and argv[i + 3].isdigit() else 1
    report: dict = {"pdf": str(pdf), "regole": str(regole), "pagina": pagina, "esito": "fallito"}
    try:
        _collauda_wizard(pdf, regole, report, pagina)
    except Exception as e:
        report["errore"] = f"{type(e).__name__}: {e}"
        report["traceback"] = traceback.format_exc()
    regole.with_suffix(".collaudo.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["esito"] == "ok" else 1


def _collauda_wizard(pdf: Path, regole: Path, report: dict, pagina: int = 1) -> None:
    from tkinter import messagebox

    import customtkinter as ctk
    from PIL import ImageGrab

    import config
    from gui.wizard_fornitore import WizardFornitore
    from processors.estrazione import estrai_listino

    messagebox.showinfo = lambda *a, **k: None  # nessuna finestra modale da chiudere a mano
    radice = ctk.CTk()
    radice.withdraw()
    w = WizardFornitore(radice, config.elenco_fornitori(), pdf=pdf, percorso_regole=regole)
    w.pagina = pagina  # prima che parta l'analisi, programmata dal costruttore
    # la pagina la sceglie chi lancia il collaudo: niente salto alla pagina consigliata, che ripartirebbe con
    # l'analisi mentre il collaudo chiede gia' l'anteprima
    w._salto_automatico = False
    w.geometry("1440x920+10+10")
    w.attributes("-topmost", True)

    def attendi(condizione, secondi=90):
        limite = time.monotonic() + secondi
        while not condizione():
            if time.monotonic() > limite:
                raise TimeoutError("tempo scaduto")
            radice.update()
            time.sleep(0.05)

    attendi(lambda: w.modello is not None)
    report["colonne"] = [[c.etichetta, campo] for c, campo in zip(w.modello.analisi.colonne, w.modello.campi, strict=True)]
    report["ancora"] = w.modello.ancora
    w._anteprima()
    attendi(lambda: w._lavoro is None)
    report["anteprima"] = w.lbl_anteprima.cget("text")
    for _ in range(10):
        radice.update()
        time.sleep(0.05)
    x, y = w.winfo_rootx(), w.winfo_rooty()
    ImageGrab.grab((x, y, x + w.winfo_width(), y + w.winfo_height())).save(regole.with_suffix(".collaudo.png"))
    w.ent_nome.insert(0, "Collaudo wizard")
    salvato = w._salva()
    radice.destroy()
    dati = config._leggi(regole)
    config.valida_schema(dati)
    profilo = dati["fornitori"]["collaudo_wizard"]
    ris = estrai_listino(pdf, profilo)
    report.update(salvato=salvato, bom=regole.read_bytes().startswith(b"\xef\xbb\xbf"), articoli=len(ris.articoli),
                  da_verificare=sum(1 for a in ris.articoli if a.tutte_le_anomalie()))
    report["esito"] = "ok" if salvato and report["bom"] and ris.articoli else "fallito"
