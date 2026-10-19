"""Scrittura del foglio "Listino" (solo le colonne del documento) e del report anomalie separato.

Le colonne arrivano dall'estrazione (`RisultatoEstrazione.colonne`): nome e ordine li decide il profilo,
il formato il ruolo. Senza colonne valgono le cinque standard, come prima delle colonne dinamiche.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from processors.models import COLONNE_STANDARD, TESTO_GENERICO, Articolo, ColonnaUscita

LARGHEZZE = {"modello": 30, "codice": 24, "ean": 16, "descrizione": 90, "prezzo": 14, TESTO_GENERICO: 24}
FORMATO_PREZZO = "#,##0.00"
# testo: niente zeri iniziali persi, niente codici letti come numeri o date
RUOLI_TESTO = {"modello", "codice", "ean", TESTO_GENERICO}


class EsportazioneFallita(Exception):
    pass


def _salva(wb: Workbook, percorso: Path) -> None:
    try:
        wb.save(percorso)
    except PermissionError as e:
        raise EsportazioneFallita(
            f"Impossibile scrivere {percorso.name}: il file e' aperto in un altro programma (Excel?)."
        ) from e
    except OSError as e:
        raise EsportazioneFallita(f"Impossibile scrivere {percorso}: {e}") from e


def _cella(a: Articolo, col: ColonnaUscita):
    if col.ruolo == "prezzo":
        return float(a.prezzo) if a.prezzo is not None else None  # mancante: cella vuota, non 0
    return a.valore(col.chiave)


def esporta_listino(articoli: list[Articolo], percorso: str | Path,
                    colonne: list[ColonnaUscita] | None = None) -> int:
    """Scrive gli articoli nel foglio "Listino" e ritorna il numero di righe esportate."""
    if not articoli:
        raise EsportazioneFallita("Nessun articolo da esportare.")
    percorso = Path(percorso)
    colonne = list(colonne or COLONNE_STANDARD)
    wb = Workbook()
    ws = wb.active
    ws.title = "Listino"
    ws.append([c.etichetta for c in colonne])
    for cella in ws[1]:
        cella.font = Font(bold=True)
        cella.alignment = Alignment(vertical="center")

    for a in articoli:
        ws.append([_cella(a, c) for c in colonne])

    for riga in ws.iter_rows(min_row=2):
        for cella, col in zip(riga, colonne, strict=True):
            if col.ruolo in RUOLI_TESTO:
                cella.number_format = "@"
            elif col.ruolo == "prezzo":
                cella.number_format = FORMATO_PREZZO
            cella.alignment = Alignment(wrap_text=col.ruolo == "descrizione", vertical="top")

    for n, col in enumerate(colonne, 1):
        ws.column_dimensions[get_column_letter(n)].width = LARGHEZZE[col.ruolo]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _salva(wb, percorso)
    return len(articoli)


def esporta_report_anomalie(articoli: list[Articolo], percorso: str | Path,
                            colonne: list[ColonnaUscita] | None = None) -> int:
    """File separato, per la verifica: non tocca la struttura del foglio "Listino".

    Ogni riga dice la pagina, le colonne che identificano l'articolo (modello e codice, se ci sono;
    altrimenti la prima colonna) e l'anomalia.
    """
    percorso = Path(percorso)
    colonne = list(colonne or COLONNE_STANDARD)
    chi = [c for c in colonne if c.ruolo in ("modello", "codice")] or colonne[:1]
    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalie"
    ws.append(["Pagina", *(c.etichetta for c in chi), "Anomalia"])
    for cella in ws[1]:
        cella.font = Font(bold=True)
    righe = 0
    for a in articoli:
        for anomalia in a.tutte_le_anomalie():
            ws.append([a.pagina, *(_cella(a, c) for c in chi), anomalia])
            righe += 1
    for n, w in enumerate((8, *(LARGHEZZE[c.ruolo] for c in chi), 90), 1):
        ws.column_dimensions[get_column_letter(n)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _salva(wb, percorso)
    return righe
