"""Colonne dinamiche: le colonne di uscita vengono dal profilo, non da un elenco fisso di cinque campi."""
import pickle
from decimal import Decimal

import pytest

from config import RegoleNonValide, regole_fornitore, valida_schema
from processors.models import COLONNE_STANDARD, Articolo, ColonnaUscita
from processors.profilo import colonne_uscita, normalizza


def _chiavi(colonne):
    return [(c.chiave, c.etichetta, c.ruolo) for c in colonne]


def test_profilo_jbl_dà_le_cinque_colonne_di_sempre():
    # tre colonne fisiche, ma la prima contiene modello, codice ed EAN impilati
    assert colonne_uscita(regole_fornitore("jbl_exhibo")) == list(COLONNE_STANDARD)
    assert [c.etichetta for c in COLONNE_STANDARD] == ["Modello", "Codice", "EAN", "Descrizione", "Prezzo"]


def test_solo_le_colonne_del_profilo_con_le_loro_etichette():
    profilo = normalizza({"nome": "Catalogo", "colonne": [
        {"campo": "ignora", "intestazioni": []},
        {"campo": "modello", "etichetta": "Articolo"},
        {"campo": "descrizione"},
        {"campo": "prezzo", "etichetta": "Listino"},
    ]})
    assert _chiavi(colonne_uscita(profilo)) == [
        ("modello", "Articolo", "modello"), ("descrizione", "Descrizione", "descrizione"),
        ("prezzo", "Listino", "prezzo"),
    ]


def test_colonne_di_testo_numerate_e_con_il_loro_nome():
    profilo = normalizza({"nome": "X", "colonne": [
        {"campo": "codice"},
        {"campo": "testo_generico", "intestazioni": ["Colour"]},
        {"campo": "testo_generico", "etichetta": "Note"},
        {"campo": "testo_generico"},
        {"campo": "codice", "etichetta": "Altro codice"},  # un ruolo standard compare una volta sola
    ]})
    assert _chiavi(colonne_uscita(profilo)) == [
        ("codice", "Codice", "codice"), ("testo_1", "Colour", "testo_generico"),
        ("testo_2", "Note", "testo_generico"), ("testo_3", "Testo 3", "testo_generico"),
    ]


def test_schema_accetta_etichetta_e_testo_e_resta_compatibile():
    vecchio = {"versione_schema": 2, "fornitori": {"a": {"nome": "A", "colonne": [{"campo": "codice"}]}}}
    valida_schema(vecchio)
    nuovo = {"versione_schema": 2, "fornitori": {"a": {"nome": "A", "colonne": [
        {"campo": "codice", "etichetta": "Cod. Fornitore"}, {"campo": "testo_generico", "intestazioni": ["Colore"]}]}}}
    valida_schema(nuovo)
    nuovo["fornitori"]["a"]["colonne"][0]["etichetta"] = ""
    with pytest.raises(RegoleNonValide):
        valida_schema(nuovo)


def test_valore_e_imposta_per_qualsiasi_colonna():
    a = Articolo("M1", "(C1)", note_estrazione={"testo_1": "Testo in piu'", "codice": "x"})
    assert (a.valore("modello"), a.valore("codice"), a.valore("testo_1")) == ("M1", "(C1)", "")
    a.imposta("testo_1", "Nero")
    assert a.valore("testo_1") == "Nero" and a.altri == {"testo_1": "Nero"} and "testo_1" not in a.note_estrazione
    a.imposta("codice", "(C2)")
    assert a.codice == "(C2)" and a.note_estrazione == {}
    a.imposta("prezzo", Decimal("5.50"))
    assert a.valore("prezzo") == Decimal("5.50")


def test_risultato_e_colonne_viaggiano_fra_processi():
    from processors.estrazione import RisultatoEstrazione

    r = RisultatoEstrazione([Articolo("M", altri={"testo_1": "Rosso"})], 1,
                            colonne=[ColonnaUscita("modello", "Articolo", "modello"),
                                     ColonnaUscita("testo_1", "Colore", "testo_generico")])
    copia = pickle.loads(pickle.dumps(r))
    assert copia.colonne == r.colonne and copia.articoli[0].altri == {"testo_1": "Rosso"}


# --- export ----------------------------------------------------------------------------------------

_TRE = [ColonnaUscita("modello", "Articolo", "modello"), ColonnaUscita("descrizione", "Descrizione", "descrizione"),
        ColonnaUscita("prezzo", "Listino", "prezzo")]


def test_excel_con_le_sole_colonne_attive(tmp_path):
    from openpyxl import load_workbook

    from exporters.excel_exporter import esporta_listino

    arts = [Articolo("VMHM4I", descrizione="Scheda input", prezzo=Decimal("888.00"), pagina=22),
            Articolo("0007", descrizione="Staffa", pagina=23)]
    colonne = [*_TRE[:2], ColonnaUscita("testo_1", "Colore", "testo_generico"), _TRE[2]]
    arts[0].altri["testo_1"] = "0123"  # testo: lo zero iniziale resta
    p = tmp_path / "tre.xlsx"
    assert esporta_listino(arts, p, colonne) == 2
    ws = load_workbook(p)["Listino"]
    assert [c.value for c in ws[1]] == ["Articolo", "Descrizione", "Colore", "Listino"]
    assert ws.max_column == 4 and ws.auto_filter.ref == "A1:D3"
    articolo, descrizione, colore, listino = ws[2]
    assert (articolo.value, colore.value, listino.value) == ("VMHM4I", "0123", 888.0)
    assert articolo.number_format == colore.number_format == "@" and listino.number_format == "#,##0.00"
    assert descrizione.alignment.wrap_text
    assert ws["A3"].value == "0007" and ws["D3"].value is None  # prezzo mancante: cella vuota


def test_report_anomalie_con_le_colonne_che_identificano(tmp_path):
    from openpyxl import load_workbook

    from exporters.excel_exporter import esporta_report_anomalie

    a = Articolo("VMHM4I", pagina=22, note_estrazione={"prezzo": "Prezzo mancante"})
    p = tmp_path / "an.xlsx"
    assert esporta_report_anomalie([a], p, _TRE) == 1
    ws = load_workbook(p)["Anomalie"]
    assert [c.value for c in ws[1]] == ["Pagina", "Articolo", "Anomalia"]
    assert [c.value for c in ws[2]] == [22, "VMHM4I", "Prezzo mancante"]
