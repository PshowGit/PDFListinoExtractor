"""Test dell'anteprima senza mostrarla: crea la finestra Tk, la usa e la distrugge."""
import time
import tkinter as tk
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from exporters.excel_exporter import esporta_listino
from processors.estrazione import estrai_listino


@pytest.fixture
def radice():
    try:
        import customtkinter as ctk

        root = ctk.CTk()
    except tk.TclError as e:
        pytest.skip(f"nessun display disponibile: {e}")
    root.withdraw()
    yield root
    root.destroy()


@pytest.fixture
def tabella(radice, pdf_sintetico, regole):
    from gui.preview_table import PreviewTable

    t = PreviewTable(radice, regole)
    t.carica(estrai_listino(pdf_sintetico, regole).articoli)
    return t


def test_righe_esportate_uguali_ai_record_approvati(tabella, tmp_path):
    ids = tabella.tree.get_children()
    assert len(ids) == 9
    tabella.tree.selection_set(ids[1:3])
    assert tabella.elimina_selezionati() == 2
    p = tmp_path / "approvati.xlsx"
    assert esporta_listino(tabella.articoli(), p) == 7
    assert load_workbook(p)["Listino"].max_row == 8  # 7 righe + intestazione


def test_il_filtro_anomalie_non_riduce_l_esportazione(tabella):
    tabella.imposta_filtro_anomalie(True)
    mostrate = len(tabella.tree.get_children())
    assert 0 < mostrate < 9
    assert len(tabella.articoli()) == 9


def test_filtro_per_tipo_di_anomalia(tabella):
    from gui.preview_table import DA_VERIFICARE, TUTTE

    tipi = dict(tabella.tipi_anomalia())
    assert tipi["Duplicato"] == 2 and tipi["Prezzo mancante"] >= 1
    tabella.imposta_filtro("Duplicato")
    assert tabella.mostrate() == 2
    assert {tabella.articolo(i).codice for i in tabella.tree.get_children()} == {"(00123-A)"}
    tabella.imposta_filtro(DA_VERIFICARE)
    assert tabella.mostrate() == tabella.conteggi()[1]
    tabella.imposta_filtro(TUTTE)
    assert tabella.mostrate() == 9
    # il conteggio riguarda tutte le righe, non solo quelle mostrate
    tabella.imposta_filtro("Prezzo mancante")
    assert dict(tabella.tipi_anomalia()) == tipi


# --- tastiera ------------------------------------------------------------------------------------

def _riga(tabella, codice):
    return next(i for i in tabella.tree.get_children() if tabella.articolo(i).codice == codice)


def test_tasti_collegati(tabella):
    for tasto in ("<Return>", "<F2>", "<Left>", "<Right>", "<F8>", "<Shift-F8>", "<Delete>"):
        assert tabella.tree.bind(tasto), tasto


def test_frecce_cambiano_la_colonna_attiva(tabella):
    assert tabella.attiva == 0 and tabella.tree.heading("modello", "text") == "▸ Modello"
    tabella.sposta_colonna(1)
    assert tabella.attiva == 1 and tabella.tree.heading("codice", "text") == "▸ Codice"
    assert tabella.tree.heading("modello", "text") == "Modello"
    tabella.sposta_colonna(-5)
    assert tabella.attiva == 0
    tabella.sposta_colonna(99)  # non oltre l'ultima colonna del documento: pagina e anomalie non si correggono
    assert tabella.attiva == len(tabella.colonne) - 1
    tabella.ordina("prezzo")  # la freccia dell'ordinamento convive col segno della colonna attiva
    assert tabella.tree.heading("prezzo", "text") == "▸ Prezzo ▲"


def test_invio_corregge_e_scende_di_riga(tabella):
    righe = tabella.tree.get_children()
    tabella.seleziona(righe[0])
    tabella.sposta_colonna(1)  # Codice
    tabella.modifica_attiva()
    assert tabella._editor is not None and tabella._editor.get() == tabella.articolo(righe[0]).codice
    tabella._editor.delete(0, "end")
    tabella._editor.insert(0, "(NUOVO-1)")
    assert tabella.conferma_editor("giu")
    assert tabella.articolo(righe[0]).codice == "(NUOVO-1)"
    assert tabella._editor is None and tabella.tree.focus() == righe[1] and tabella.tree.selection() == (righe[1],)


def test_tab_apre_la_cella_accanto_ed_esc_annulla(tabella):
    righe = tabella.tree.get_children()
    tabella.modifica(righe[0], 1)  # Codice
    prima = tabella.articolo(righe[0]).codice
    assert tabella.conferma_editor("destra")
    assert tabella.attiva == 2 and tabella._in_modifica == (righe[0], 2)  # EAN, stessa riga
    tabella._editor.delete(0, "end")
    tabella._editor.insert(0, "123")
    tabella.annulla_editor()
    assert tabella._editor is None and tabella.articolo(righe[0]).ean != "123"
    assert tabella.articolo(righe[0]).codice == prima  # confermato senza cambiarlo
    tabella.modifica(righe[0], 1)
    assert tabella.conferma_editor("sinistra") and tabella._in_modifica == (righe[0], 0)
    tabella.annulla_editor()


def test_col_filtro_la_riga_corretta_lascia_il_posto_alla_successiva(tabella):
    from gui.preview_table import DA_VERIFICARE

    tabella.imposta_filtro(DA_VERIFICARE)
    da_verificare = list(tabella.tree.get_children())
    iid = _riga(tabella, "(SPK-003)")
    dopo = da_verificare[da_verificare.index(iid) + 1]
    tabella.modifica(iid, tabella.colonne.index(next(c for c in tabella.colonne if c.ruolo == "prezzo")))
    tabella._editor.delete(0, "end")
    tabella._editor.insert(0, "1.090,50")
    assert tabella.conferma_editor("giu")
    assert not tabella.tree.exists(iid)  # corretta: non e' piu' da verificare
    assert tabella.tree.focus() == dopo and tabella.tree.selection() == (dopo,)  # non si scende di due


def test_prezzo_non_valido_lascia_l_editor_aperto(tabella, monkeypatch):
    from tkinter import messagebox

    monkeypatch.setattr(messagebox, "showerror", lambda *a, **k: None)
    iid = _riga(tabella, "(SPK-003)")
    tabella.modifica(iid, next(n for n, c in enumerate(tabella.colonne) if c.ruolo == "prezzo"))
    tabella._editor.delete(0, "end")
    tabella._editor.insert(0, "boh")
    assert not tabella.conferma_editor("giu")
    assert tabella._editor is not None and tabella._in_modifica[0] == iid
    tabella.annulla_editor()


def test_f8_va_alle_righe_da_verificare(tabella):
    righe = list(tabella.tree.get_children())
    con = [i for i in righe if tabella.articolo(i).tutte_le_anomalie()]
    assert len(con) >= 2
    tabella.seleziona(righe[0] if righe[0] not in con else righe[righe.index(con[0]) + 1])
    primo = tabella.vai_da_verificare(1)
    assert primo in con and tabella.tree.selection() == (primo,)
    secondo = tabella.vai_da_verificare(1)
    assert secondo == con[(con.index(primo) + 1) % len(con)]
    assert tabella.vai_da_verificare(-1) == primo
    tabella.seleziona(con[-1])
    assert tabella.vai_da_verificare(1) == con[0]  # ricomincia dall'inizio


def test_eliminare_la_riga_selezionata_seleziona_la_successiva(tabella):
    righe = list(tabella.tree.get_children())
    tabella.seleziona(righe[2])
    tabella.elimina_selezionati()
    assert tabella.tree.selection() == (righe[3],)


def test_correggere_il_prezzo_toglie_l_anomalia(tabella):
    iid = next(i for i in tabella.tree.get_children() if tabella.articolo(i).codice == "(SPK-003)")
    assert tabella.articolo(iid).tutte_le_anomalie()
    assert tabella.modifica_campo(iid, "prezzo", "1.090,50")
    a = tabella.articolo(iid)
    assert a.prezzo == Decimal("1090.50") and not a.tutte_le_anomalie()
    assert not tabella.modifica_campo(iid, "prezzo", "boh")
    assert a.prezzo == Decimal("1090.50")


def test_eliminare_un_duplicato_aggiorna_l_altro(tabella):
    dup = [i for i in tabella.tree.get_children() if tabella.articolo(i).codice == "(00123-A)"]
    assert len(dup) == 2
    tabella.tree.selection_set(dup[1])
    tabella.elimina_selezionati()
    assert not any("Duplicato" in x for x in tabella.articolo(dup[0]).tutte_le_anomalie())


def test_ordinamento_per_prezzo_con_mancanti_in_fondo(tabella):
    tabella.ordina("prezzo")
    prezzi = [a.prezzo for a in tabella.articoli()]
    validi = [p for p in prezzi if p is not None]
    assert validi == sorted(validi) and prezzi[-1] is None
    tabella.ordina("prezzo")
    prezzi = [a.prezzo for a in tabella.articoli()]
    assert [p for p in prezzi if p is not None] == sorted(validi, reverse=True) and prezzi[-1] is None


def _attendi(radice, condizione, secondi=60):
    import time

    limite = time.monotonic() + secondi
    while not condizione():
        if time.monotonic() > limite:
            raise AssertionError("tempo scaduto")
        radice.update()
        time.sleep(0.05)


def test_wizard_dalla_pdf_al_fornitore_salvato(radice, tmp_path, monkeypatch):
    from tkinter import messagebox

    import config
    from gui.wizard_fornitore import WizardFornitore
    from tests.sintetico_generico import crea_listino_generico

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setattr(messagebox, "showinfo", lambda *a, **k: None)
    pdf = crea_listino_generico(tmp_path / "Listino ACME è.pdf")
    salvati = []
    w = WizardFornitore(radice, config.elenco_fornitori(), al_salvataggio=salvati.append, pdf=pdf)
    w.withdraw()
    _attendi(radice, lambda: w.modello is not None)  # analisi nel processo separato
    assert [m.get() for m in w._menu_colonne] == ["Codice", "Descrizione", "EAN", "Prezzo"]
    assert [n.get() for n in w._nomi_colonne] == ["SKU", "Description", "EAN", "Price EUR"]
    assert w.lbl_totale.cget("text") == "di 3"

    # un cambio di mappatura arriva al modello e si vede sulla pagina; il nome segue il ruolo
    w.seleziona_colonna(2)
    w.imposta_campo(2, "ignora")
    assert w.modello.campi[2] == "ignora" and w._nomi_colonne[2].get() == ""
    w.imposta_campo(2, "ean")
    assert w._nomi_colonne[2].get() == "EAN"

    w._anteprima()
    _attendi(radice, lambda: w._lavoro is None)
    assert [a.codice for a in w.tabella.articoli()] == ["SP-100", "SP-200", "SP-300", "0007-X", "AMP-1"]
    assert "5 articoli" in w.lbl_anteprima.cget("text")

    w.ent_nome.insert(0, "ACME Audio")
    assert w._salva()
    assert salvati == ["acme_audio"]
    assert (tmp_path / "PDFListinoExtractor" / "supplier_rules.json").read_bytes().startswith(b"\xef\xbb\xbf")
    assert config.elenco_fornitori()["acme_audio"] == "ACME Audio"


def test_wizard_in_primo_piano_con_le_colonne_compatte(radice, tmp_path):
    from gui.wizard_fornitore import WizardFornitore
    from tests.sintetico_generico import crea_listino_colonne

    w = WizardFornitore(radice, {}, pdf=crea_listino_colonne(tmp_path / "prezzi.pdf", "prezzi"))
    try:
        # davanti alla finestra principale: su Windows un CTkToplevel appena aperto finiva dietro
        assert w.attributes("-topmost") == 1
        _attendi(radice, lambda: w.modello is not None)
        assert w.lbl_colonne.cget("text") == "Colonne rilevate (5)"
        assert [m.get() for m in w._menu_colonne] == [
            "Codice", "Descrizione", "Prezzo", "Altro prezzo (colonna in piu')", "Altro prezzo (colonna in piu')"]
        # ancora, formato e parentesi richiusi, coi valori riassunti; si aprono da soli se l'errore e' li'
        assert not w._opzioni_aperte and not w._riquadro_opzioni.grid_info()
        assert "Altre opzioni: ancora Codice" in w.btn_opzioni.cget("text")
        w.modello.ancora = "ean"
        w._mostra_errori()
        assert w._opzioni_aperte and w._riquadro_opzioni.grid_info()
        w.mostra_opzioni_avanzate(False)
        assert not w._riquadro_opzioni.grid_info() and w.btn_opzioni.cget("text").startswith("▸")
    finally:
        w.chiudi()


def test_wizard_disegna_l_anteprima_sulla_pagina(radice, tmp_path):
    from gui.wizard_fornitore import WizardFornitore
    from tests.sintetico_catalogo import crea_catalogo

    pdf = crea_catalogo(tmp_path / "catalogo.pdf")
    w = WizardFornitore(radice, {}, pdf=pdf)
    w.withdraw()
    try:
        _attendi(radice, lambda: w.modello is not None)
        assert w.modello.analisi.intestazioni and len(w.modello.analisi.intestazioni) == 2  # due tabelle
        w._vai(2)  # la pagina col titolo «CAVI» e il bollino «NEW»
        _attendi(radice, lambda: w._lavoro is None and w.modello.analisi.pagina == 2)
        w._anteprima()
        _attendi(radice, lambda: w._lavoro is None)
        articoli = w.tabella.articoli()
        # pagina e mappatura affiancate, anteprima sotto: separatori trascinabili
        assert len(w._corpo.panes()) == 2 and len(w._sopra.panes()) == 2
        # un riquadro per articolo, coi titoli e i bollini visti dal motore
        assert [a for a, _ in w.riquadri_articoli] == articoli and len(articoli) == 3
        assert [t[4] for t in w._geometria.titoli] == ["CAVI"] and [b[4] for b in w._geometria.bollini] == ["NEW"]
        assert "1 titoli, 1 bollini" in w.lbl_anteprima.cget("text")
        # dalla tabella alla pagina e dalla pagina alla tabella
        terzo = articoli[2]
        w._riga_anteprima(terzo)
        assert w._evidenziato is terzo
        x0, y0, x1, y1 = terzo.bbox
        assert w.articolo_nel_punto((x0 + x1) / 2, (y0 + y1) / 2) is terzo
        assert w.articolo_nel_punto(5, 5) is None
        # cambiare la mappatura toglie i riquadri: venivano dalla mappatura di prima
        w.imposta_campo(2, "ignora")
        assert w.riquadri_articoli == [] and "da rifare" in w.lbl_anteprima.cget("text")
    finally:
        w.chiudi()


def test_wizard_salta_l_indice_e_segna_le_pagine(radice, tmp_path):
    from gui.wizard_fornitore import WizardFornitore
    from tests.sintetico_catalogo import crea_catalogo_con_indice

    w = WizardFornitore(radice, {}, pdf=crea_catalogo_con_indice(tmp_path / "con_indice.pdf"))
    w.withdraw()
    try:
        # la pagina 1 e' l'indice: il wizard passa da solo alla pagina consigliata e lo dice
        _attendi(radice, lambda: w.modello is not None and w.modello.analisi.pagina == 2 and w._lavoro is None)
        assert w.pagina == 2 and w.modello.campi == ["marca", "modello", "descrizione", "prezzo"]
        assert "La pagina 1 sembra un indice" in w.lbl_stato.cget("text")
        voci = w.menu_campione.cget("values")
        assert voci[0] == "○ pag. 1 · indice o divisorio" and w.menu_campione.get().startswith("✔ pag. 2")
        # la prossima tabella; e tornando all'indice a mano non si salta piu', si suggerisce
        w._prossima_tabella()
        _attendi(radice, lambda: w._lavoro is None and w.modello.analisi.pagina == 3)
        w._scegli_pagina_campione(voci[0])
        _attendi(radice, lambda: w._lavoro is None and w.modello.analisi.pagina == 1)
        assert w.pagina == 1 and "Prova la pagina 2" in w.lbl_stato.cget("text")
    finally:
        w.chiudi()


def test_wizard_con_pdf_non_valido_mostra_l_errore(radice, tmp_path, monkeypatch):
    from tkinter import messagebox

    from gui.wizard_fornitore import WizardFornitore

    errori = []
    monkeypatch.setattr(messagebox, "showerror", lambda titolo, testo, **k: errori.append(testo))
    rotto = tmp_path / "rotto.pdf"
    rotto.write_bytes(b"x" * 50)
    w = WizardFornitore(radice, {}, pdf=rotto)
    w.withdraw()
    _attendi(radice, lambda: errori)
    assert "PDF" in errori[0] and w.modello is None
    w.chiudi()


def test_il_fornitore_salvato_compare_e_si_seleziona_nella_finestra_principale(radice, tmp_path, monkeypatch):
    from config import salva_fornitore_personale
    from gui.main_window import FinestraPrincipale

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    radice.destroy()
    app = FinestraPrincipale()
    app.withdraw()
    try:
        salva_fornitore_personale("acme", {"nome": "ACME", "colonne": [{"campo": "codice"}, {"campo": "prezzo"}]})
        app._fornitore_salvato("acme")
        assert app.menu_fornitore.get() == "ACME" and app.chiave_fornitore == "acme"
        assert "ACME" in app.menu_fornitore.cget("values")
        assert app.regole["colonne"][0]["campo"] == "codice"
    finally:
        app.destroy()
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def test_la_finestra_principale_si_costruisce(radice):
    from gui.main_window import FinestraPrincipale

    radice.destroy()
    app = FinestraPrincipale()
    app.withdraw()
    try:
        assert app.btn_estrai.cget("state") == "disabled"
        assert app.btn_esporta.cget("state") == "disabled"
    finally:
        app.destroy()
    # la fixture distruggera' una radice gia' distrutta: si ricrea per simmetria
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def test_il_dialogo_del_nome_parte_dal_nome_attuale(radice):
    from gui.dialoghi import DialogoTesto

    d = DialogoTesto(radice, "Rinomina fornitore", "Nuovo nome:", "ACME")
    assert d.campo.get() == "ACME"
    d.campo.delete(0, "end")
    d.campo.insert(0, "ACME Audio")
    d._conferma()  # come Invio o il pulsante OK
    assert d.risultato == "ACME Audio" and not d.winfo_exists()

    d = DialogoTesto(radice, "Rinomina fornitore", "Nuovo nome:", "ACME")
    d.destroy()  # come Esc, Annulla o la chiusura della finestra
    assert d.risultato is None and not d.winfo_exists()


def test_rinominare_e_rimuovere_un_fornitore_dalla_finestra_principale(radice, tmp_path, monkeypatch):
    from tkinter import messagebox

    import gui.main_window as finestra
    from config import salva_fornitore_personale

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    salva_fornitore_personale("acme", {"nome": "ACME", "colonne": [{"campo": "codice"}, {"campo": "prezzo"}]})
    avvisi = []
    monkeypatch.setattr(messagebox, "showwarning", lambda titolo, testo, **k: avvisi.append(testo))
    radice.destroy()
    app = finestra.FinestraPrincipale()
    app.withdraw()
    try:
        distribuito = app.fornitori["jbl_exhibo"]
        # il fornitore distribuito col programma e' in sola lettura
        assert app.chiave_fornitore == "jbl_exhibo"
        assert app.btn_rinomina.cget("state") == "disabled" and app.btn_rimuovi.cget("state") == "disabled"
        assert app.btn_modifica.cget("state") == "disabled"
        app._cambia_fornitore("ACME")
        assert app.btn_rinomina.cget("state") == "normal" and app.btn_rimuovi.cget("state") == "normal"
        assert app.btn_modifica.cget("state") == "normal"

        monkeypatch.setattr(finestra, "chiedi_testo", lambda *a, **k: "ACME Audio 2026")
        app._rinomina_fornitore()
        assert app.menu_fornitore.get() == "ACME Audio 2026" and app.chiave_fornitore == "acme"
        assert list(app.menu_fornitore.cget("values")) == [distribuito, "ACME Audio 2026"]
        assert app.regole["nome"] == "ACME Audio 2026"

        monkeypatch.setattr(finestra, "chiedi_testo", lambda *a, **k: distribuito.lower())
        app._rinomina_fornitore()
        assert "Esiste gia'" in avvisi[0] and app.menu_fornitore.get() == "ACME Audio 2026"

        monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: False)
        app._rimuovi_fornitore()
        assert "acme" in app.fornitori

        monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: True)
        app._rimuovi_fornitore()
        assert list(app.fornitori) == ["jbl_exhibo"] and app.chiave_fornitore == "jbl_exhibo"
        assert list(app.menu_fornitore.cget("values")) == [distribuito]
        assert app.menu_fornitore.get() == distribuito
        assert app.btn_rimuovi.cget("state") == "disabled"
    finally:
        app.destroy()
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def test_esportare_e_importare_i_fornitori_dalla_finestra_principale(radice, tmp_path, monkeypatch):
    import json
    from tkinter import filedialog, messagebox

    import gui.main_window as finestra
    from config import salva_fornitore_personale

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "vecchia"))
    salva_fornitore_personale("acme", {"nome": "ACME", "colonne": [{"campo": "codice"}, {"campo": "prezzo"}]})
    esportato = tmp_path / "Fornitori.json"
    messaggi = []
    monkeypatch.setattr(messagebox, "showinfo", lambda titolo, testo, **k: messaggi.append(testo))
    monkeypatch.setattr(messagebox, "showerror", lambda titolo, testo, **k: messaggi.append(("errore", testo)))
    monkeypatch.setattr(filedialog, "asksaveasfilename", lambda **k: str(esportato))
    monkeypatch.setattr(filedialog, "askopenfilename", lambda **k: str(esportato))
    radice.destroy()
    app = finestra.FinestraPrincipale()
    app.withdraw()
    try:
        assert app.btn_esporta_fornitori.cget("state") == "normal"
        app._esporta_fornitori()
        assert esportato.is_file() and "Esportato 1 fornitore" in messaggi[-1]
    finally:
        app.destroy()

    # la postazione nuova ha gia' un suo «ACME»: quello importato non lo sovrascrive
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "nuova"))
    salva_fornitore_personale("acme", {"nome": "ACME", "colonne": [{"campo": "ean"}, {"campo": "prezzo"}]})
    app = finestra.FinestraPrincipale()
    app.withdraw()
    try:
        distribuito = app.fornitori["jbl_exhibo"]
        app._importa_fornitori()
        assert list(app.menu_fornitore.cget("values")) == [distribuito, "ACME", "ACME - Importato"]
        assert app.chiave_fornitore == "acme_importato" and app.menu_fornitore.get() == "ACME - Importato"
        assert app.regole["colonne"][0]["campo"] == "codice"
        assert "acme_importato" in app.personali and app.btn_rimuovi.cget("state") == "normal"
        assert "«ACME» → «ACME - Importato»" in messaggi[-1]

        esportato.write_text(json.dumps({"versione_schema": 2, "fornitori": {"x": {"nome": "X"}}}), encoding="utf-8")
        app._importa_fornitori()
        assert messaggi[-1][0] == "errore" and "Non e' stato importato nulla" in messaggi[-1][1]
        assert list(app.fornitori) == ["jbl_exhibo", "acme", "acme_importato"]
    finally:
        app.destroy()
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def test_wizard_in_modifica_riprende_e_aggiorna_il_fornitore(radice, tmp_path, monkeypatch):
    from tkinter import messagebox

    import config
    from gui.wizard_fornitore import WizardFornitore
    from processors.analisi_layout import analizza_pdf
    from processors.wizard import ModelloWizard
    from tests.sintetico_generico import crea_listino_generico

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setattr(messagebox, "showinfo", lambda *a, **k: None)
    pdf = crea_listino_generico(tmp_path / "acme.pdf")
    m = ModelloWizard(analizza_pdf(str(pdf), 1))
    m.nome, m.formato_prezzo = "ACME", "it"
    m.imposta_campo(2, "ignora")
    config.salva_fornitore_personale("acme", {**m.profilo(), "escludi": ["PROMO"]})

    salvati = []
    w = WizardFornitore(radice, config.elenco_fornitori(), al_salvataggio=salvati.append, pdf=pdf,
                        modifica=("acme", config.profilo_personale("acme")))
    w.withdraw()
    _attendi(radice, lambda: w.modello is not None)
    assert "Modifica fornitore" in w.title() and w.btn_salva.cget("text") == "Salva modifiche"
    assert [x.get() for x in w._menu_colonne] == ["Codice", "Descrizione", "— ignora —", "Prezzo"]
    assert [x.get() for x in w._nomi_colonne] == ["SKU", "Description", "", "Price EUR"]
    assert w.ent_nome.get() == "ACME" and w.menu_formato.get() == "Italiano (1.234,56)"
    assert "4 colonne su 4 riprese" in w.lbl_stato.cget("text")

    w.imposta_campo(2, "ean")
    w.ent_nome.delete(0, "end")
    w.ent_nome.insert(0, "ACME Audio")
    assert w._salva()
    assert salvati == ["acme"]
    salvato = config.profilo_personale("acme")
    assert salvato["nome"] == "ACME Audio" and salvato["colonne"][2]["campo"] == "ean"
    assert salvato["escludi"] == ["PROMO"]  # cio' che il wizard non mostra resta
    assert set(config.chiavi_personali()) == {"acme"}


def test_tabella_con_le_colonne_dell_estrazione(radice):
    from gui.preview_table import PreviewTable
    from processors.models import Articolo, ColonnaUscita

    colonne = [ColonnaUscita("modello", "Articolo", "modello"), ColonnaUscita("testo_1", "Colore", "testo_generico"),
               ColonnaUscita("prezzo", "Listino", "prezzo")]
    regole = {"nome": "X", "colonne": [{"campo": "modello"}, {"campo": "testo_generico"}, {"campo": "prezzo"}]}
    t = PreviewTable(radice, regole)
    arts = [Articolo("B-2", prezzo=Decimal("5.00"), pagina=1, altri={"testo_1": "Rosso"}),
            Articolo("A-1", pagina=1, note_estrazione={"testo_1": "Testo in piu' a pagina 2"}),
            Articolo("C-3", prezzo=Decimal("2.50"), pagina=2)]
    t.carica(arts, colonne)
    assert t.tree["columns"] == ("modello", "testo_1", "prezzo", "pagina", "note")
    # «▸» segna la colonna attiva per la tastiera (la prima)
    assert [t.tree.heading(c, "text") for c in ("modello", "testo_1", "prezzo")] == ["▸ Articolo", "Colore", "Listino"]
    assert t.tree.item(t.tree.get_children()[0], "values")[:3] == ("B-2", "Rosso", "5,00")
    for _ in range(2):  # crescente e decrescente: il prezzo mancante resta in fondo
        t.ordina("prezzo")
        assert t.articoli()[-1].modello == "A-1"
    t.ordina("testo_1")
    assert [a.modello for a in t.articoli()][-1] == "B-2"  # "" < "rosso"
    iid = next(i for i in t.tree.get_children() if t.articolo(i).modello == "A-1")
    assert t.modifica_campo(iid, "testo_1", " Blu ")
    a = t.articolo(iid)
    assert a.altri["testo_1"] == "Blu" and "testo_1" not in a.note_estrazione
    assert a.anomalie == ["Prezzo mancante"]  # anomalie ricalcolate: il profilo prevede il prezzo


@pytest.mark.parametrize("variante", ["tre", "colore", "prezzi"])
def test_colonne_dinamiche_dal_pdf_all_excel(radice, tmp_path, variante):
    # tutto il percorso: proposta del wizard, profilo, estrazione, tabella, foglio "Listino"
    from exporters.excel_exporter import esporta_listino
    from gui.preview_table import PreviewTable
    from processors.analisi_layout import analizza_pdf
    from processors.wizard import ModelloWizard
    from tests.sintetico_generico import ATTESI_COLONNE, crea_listino_colonne

    intestazioni, righe = ATTESI_COLONNE[variante]
    pdf = crea_listino_colonne(tmp_path / f"{variante}.pdf", variante)
    m = ModelloWizard(analizza_pdf(str(pdf), 1))
    m.nome = "ACME"
    assert m.errori() == []
    profilo = m.profilo()
    ris = estrai_listino(pdf, profilo)
    assert [c.etichetta for c in ris.colonne] == intestazioni

    t = PreviewTable(radice, profilo)
    t.carica(ris.articoli, ris.colonne)
    titoli = [t.tree.heading(c, "text").removeprefix("▸ ") for c in t.tree["columns"]]  # «▸»: colonna attiva
    assert titoli == [*intestazioni, "Pag.", "Anomalie"]
    assert not any(a.tutte_le_anomalie() for a in t.articoli())

    p = tmp_path / f"{variante}.xlsx"
    assert esporta_listino(t.articoli(), p, t.colonne) == len(righe)
    ws = load_workbook(p)["Listino"]
    assert [c.value for c in ws[1]] == intestazioni and ws.max_column == len(intestazioni)
    assert [list(r) for r in ws.iter_rows(min_row=2, values_only=True)] == righe
    if variante == "tre":
        assert not {"Codice", "EAN"} & set(intestazioni)
    elif variante == "colore":
        assert ws.cell(2, intestazioni.index("Colour") + 1).number_format == "@"
    else:
        # i prezzi in piu' sono numeri come il prezzo, e la tabella li mostra all'italiana
        assert ws.cell(2, intestazioni.index("Dealer") + 1).number_format == "#,##0.00"
        assert t._valori(t.articoli()[1])[3] == "949,00"


def test_prezzo_in_piu_si_corregge_e_si_ordina_come_un_prezzo(radice):
    from gui.preview_table import PreviewTable
    from processors.models import Articolo, ColonnaUscita

    colonne = [ColonnaUscita("codice", "Codice", "codice"), ColonnaUscita("prezzo_1", "Dealer", "altro_prezzo")]
    profilo = {"formato_prezzo": "it"}
    articoli = [Articolo(codice="A", altri={"prezzo_1": Decimal("90.00")}),
                Articolo(codice="B", altri={"prezzo_1": "Chiedere"}, note_estrazione={"prezzo_1": "Prezzo non riconosciuto"}),
                Articolo(codice="C", altri={"prezzo_1": Decimal("1200.00")})]
    t = PreviewTable(radice, profilo)
    t.carica(articoli, colonne)
    assert [t._valori(a)[1] for a in articoli] == ["90,00", "Chiedere", "1.200,00"]
    t.ordina("prezzo_1")
    assert [a.codice for a in t.articoli()] == ["A", "C", "B"]  # quello non riconosciuto in fondo
    t.ordina("prezzo_1")
    assert [a.codice for a in t.articoli()] == ["C", "A", "B"]
    iid = t.iid_di(articoli[1])
    assert not t.modifica_campo(iid, "prezzo_1", "a richiesta")
    assert t.modifica_campo(iid, "prezzo_1", "1.050,50")
    assert articoli[1].altri["prezzo_1"] == Decimal("1050.50") and "prezzo_1" not in articoli[1].note_estrazione
    assert t.modifica_campo(iid, "prezzo_1", "")
    assert articoli[1].altri["prezzo_1"] == ""


def test_piu_prezzi_trovati_si_scelgono_da_un_menu(radice):
    from tkinter import ttk

    from gui.preview_table import PreviewTable
    from processors.models import Articolo, ColonnaUscita

    a = Articolo(codice="ACT-6", prezzo_testo="€ 250,00 | € 290,00",
                 note_estrazione={"prezzo": "Piu' prezzi per lo stesso articolo (€ 250,00 | € 290,00)"})
    t = PreviewTable(radice, {"formato_prezzo": "it"})
    t.carica([a], [ColonnaUscita("codice", "Codice", "codice"), ColonnaUscita("prezzo", "Listino", "prezzo")])
    t.modifica(t.iid_di(a), 1)
    assert isinstance(t._editor, ttk.Combobox)
    assert list(t._editor.cget("values")) == ["€ 250,00", "€ 290,00"] and t._editor.get() == "€ 250,00"
    t._editor.set("€ 290,00")
    assert t.conferma_editor()
    assert a.prezzo == Decimal("290.00") and "prezzo" not in a.note_estrazione


def test_menu_del_filtro_coi_contatori(radice, tmp_path, monkeypatch, pdf_sintetico, regole):
    from gui.main_window import FinestraPrincipale

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    radice.destroy()
    app = FinestraPrincipale()
    app.withdraw()
    try:
        app.tabella.carica(estrai_listino(pdf_sintetico, regole).articoli)
        app._aggiorna_contatori()
        n, anomali = app.tabella.conteggi()
        voci = app.menu_filtro.cget("values")
        assert voci[:2] == [f"Tutte le righe ({n})", f"Da verificare ({anomali})"]
        assert "Duplicato (2)" in voci
        app._filtro("Duplicato (2)")
        assert app.tabella.mostrate() == 2 and app.menu_filtro.get() == "Duplicato (2)"
        assert app.lbl_conteggi.cget("text") == f"2 di {n} righe"
        # eliminare un duplicato aggiorna i contatori; il tipo scelto resta nel menu anche a zero
        app.tabella.tree.selection_set(app.tabella.tree.get_children()[0])
        app.tabella.elimina_selezionati()
        assert app.menu_filtro.get() == "Duplicato (0)" and app.tabella.mostrate() == 0
        assert f"Da verificare ({anomali - 2})" in app.menu_filtro.cget("values")
    finally:
        app.destroy()
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def test_avanzamento_avvisi_e_pannelli_della_finestra_principale(radice, tmp_path, monkeypatch, pdf_sintetico, regole):
    from gui.dialoghi import FinestraAvvisi
    from gui.main_window import PESO_LETTURA, FinestraPrincipale, _durata
    from processors.estrazione import RisultatoEstrazione

    assert (_durata(4.4), _durata(59.6), _durata(125)) == ("4 s", "1 min 00 s", "2 min 05 s")
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    radice.destroy()
    app = FinestraPrincipale()
    app.withdraw()
    try:
        # tabella e dettaglio in un pannello diviso da un separatore trascinabile
        assert len(app.pannello.panes()) == 2
        assert app.btn_avvisi.cget("state") == "disabled"
        # la barra si muove gia' durante la lettura (prima passata), poi nell'analisi
        app._inizio = __import__("time").monotonic()
        app._lettura(50, 100)
        assert app.barra.get() == pytest.approx(PESO_LETTURA / 2, abs=0.01)
        assert app.lbl_stato.cget("text").startswith("Lettura pagina 50 di 100 · ")
        app._progresso(100, 100, 673)
        assert app.barra.get() == pytest.approx(1, abs=0.01)
        assert app.lbl_stato.cget("text").startswith("Analisi pagina 100 di 100 · 673 articoli · ")
        # tutti gli avvisi, in una finestra che si riapre dal pulsante
        avvisi = [f"Pagina {n}: testo non attribuito" for n in range(1, 31)]
        ris = RisultatoEstrazione(estrai_listino(pdf_sintetico, regole).articoli, 4, avvisi=avvisi)
        app._estrazione_finita(ris)
        assert isinstance(app._finestra_avvisi, FinestraAvvisi) and app._finestra_avvisi.avvisi == avvisi
        assert app._finestra_avvisi.testo.get("1.0", "end").strip().splitlines() == avvisi  # tutti, non 20
        assert app.btn_avvisi.cget("text") == "Avvisi (30)…" and app.btn_avvisi.cget("state") == "normal"
        assert app.lbl_stato.cget("text").startswith("Completato: 4 pagine · 9 articoli · ")
        app._finestra_avvisi.destroy()
        app._mostra_avvisi()
        assert app._finestra_avvisi.winfo_exists()
    finally:
        app.destroy()
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def test_larghezza_delle_colonne_dal_contenuto(radice):
    from gui.preview_table import LARGHEZZA_MASSIMA, PreviewTable
    from processors.models import Articolo

    t = PreviewTable(radice, {"nome": "X", "colonne": [{"campo": "modello"}, {"campo": "descrizione"}]})
    t.carica([Articolo("A1", descrizione="x")])
    stretta = t.tree.column("modello", "width")
    t.carica([Articolo("MODELLO-CON-UN-NOME-MOLTO-LUNGO-2026", descrizione="x")])
    assert t.tree.column("modello", "width") > stretta
    t.carica([Articolo("X" * 300, descrizione="x")])
    assert t.tree.column("modello", "width") == LARGHEZZA_MASSIMA


def test_dettaglio_con_le_colonne_del_documento(radice, tmp_path, monkeypatch):
    from gui.main_window import FinestraPrincipale
    from processors.models import Articolo, ColonnaUscita

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    radice.destroy()
    app = FinestraPrincipale()
    app.withdraw()
    try:
        colonne = [ColonnaUscita("modello", "Articolo", "modello"), ColonnaUscita("descrizione", "Descrizione", "descrizione"),
                   ColonnaUscita("prezzo", "Listino", "prezzo")]
        a = Articolo("VMHM4I", descrizione="Scheda input 4 canali HDMI", prezzo=Decimal("888.00"), pagina=22,
                     categoria="VISSONIC · MODULI INPUT")
        app.tabella.carica([a], colonne)
        app._mostra_dettaglio(a)
        testo = app.dettaglio.get("1.0", "end")
        assert "Articolo VMHM4I" in testo and "888,00 €" in testo and "Scheda input 4 canali HDMI" in testo
        assert "EAN" not in testo and "Codice" not in testo
        senza = Articolo("VMX-1", pagina=22)
        app._mostra_dettaglio(senza)
        testo = app.dettaglio.get("1.0", "end")
        assert "senza prezzo" in testo and "€" not in testo
    finally:
        app.destroy()
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def _finestra_con_aggiornamento(radice, tmp_path, monkeypatch, rilascio):
    """Finestra principale come nell'eseguibile installato, con l'API di GitHub simulata."""
    import aggiornamento
    import gui.main_window as finestra

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setattr(aggiornamento, "attivo", lambda: True)
    monkeypatch.setattr(aggiornamento, "pulisci_scaricati", lambda: None)
    monkeypatch.setattr(aggiornamento, "cerca_aggiornamento", lambda: rilascio)
    radice.destroy()
    app = finestra.FinestraPrincipale()
    app.withdraw()
    app._controlla_aggiornamento()  # all'avvio parte dopo 1,5 s
    limite = time.monotonic() + 5
    while app.rilascio is None and rilascio is not None and time.monotonic() < limite:
        app.update()
        time.sleep(0.05)
    return app, finestra


def _radice_nuova(radice):
    import customtkinter as ctk

    nuova = ctk.CTk()
    nuova.withdraw()
    radice.destroy = nuova.destroy


def _rilascio_di_prova():
    from aggiornamento import Rilascio

    return Rilascio("9.1.0", "v9.1.0", "## Novità\n\n- **Una** cosa", "https://example.com/release", "u", "0" * 64, 1)


def test_senza_aggiornamenti_il_pulsante_non_compare(radice, tmp_path, monkeypatch):
    app, _ = _finestra_con_aggiornamento(radice, tmp_path, monkeypatch, None)
    try:
        for _ in range(10):
            app.update()
            time.sleep(0.03)
        assert app.rilascio is None and not app.btn_aggiornamento.grid_info()
    finally:
        app.destroy()
    _radice_nuova(radice)


def test_aggiornamento_disponibile_e_ignorato(radice, tmp_path, monkeypatch):
    import aggiornamento

    app, finestra = _finestra_con_aggiornamento(radice, tmp_path, monkeypatch, _rilascio_di_prova())
    try:
        assert app.btn_aggiornamento.grid_info() and app.btn_aggiornamento.cget("text") == (
            "Aggiornamento disponibile (v9.1.0)")

        class Ignora(finestra.DialogoAggiornamento):
            def __init__(self, *a, **k):
                super().__init__(*a, **k)
                assert "• Una cosa" in self.note.get("1.0", "end")
                self.after(100, lambda: self._scegli("ignora"))

        monkeypatch.setattr(finestra, "DialogoAggiornamento", Ignora)
        app._apri_aggiornamento()
        assert aggiornamento.versione_ignorata() == "9.1.0" and not app.btn_aggiornamento.grid_info()
        assert not (tmp_path / "PDFListinoExtractor" / "supplier_rules.json").exists()
        # al prossimo avvio la versione ignorata non si propone piu'
        app.rilascio = None
        app._mostra_aggiornamento(_rilascio_di_prova())
        assert app.rilascio is None and not app.btn_aggiornamento.grid_info()
    finally:
        app.destroy()
    _radice_nuova(radice)


def test_aggiorna_ora_scarica_avvia_e_chiude(radice, tmp_path, monkeypatch):
    import aggiornamento

    app, finestra = _finestra_con_aggiornamento(radice, tmp_path, monkeypatch, _rilascio_di_prova())
    setup = tmp_path / "PDFListinoExtractor_Setup_9.1.0.exe"
    avviati, chiusa = [], []

    class Aggiorna(finestra.DialogoAggiornamento):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self.after(100, lambda: self._scegli("aggiorna"))

    def scarica(rilascio, avanzamento, annullato):
        for letti in (10, 20):
            avanzamento(letti, 20)
            time.sleep(0.05)
        return setup

    monkeypatch.setattr(finestra, "DialogoAggiornamento", Aggiorna)
    monkeypatch.setattr(aggiornamento, "modo_installazione", lambda: "/CURRENTUSER")
    monkeypatch.setattr(aggiornamento, "scarica", scarica)
    monkeypatch.setattr(aggiornamento, "avvia_installazione", lambda percorso, modo: avviati.append((percorso, modo)))
    monkeypatch.setattr(app, "_chiudi", lambda: chiusa.append(True))
    try:
        app._apri_aggiornamento()
        assert avviati == [(setup, "/CURRENTUSER")] and chiusa == [True]

        # file non integro: si dice perche', non si installa e il programma resta aperto
        def non_valido(*a):
            raise aggiornamento.AggiornamentoNonValido("SHA-256 diverso")

        errori = []
        monkeypatch.setattr(aggiornamento, "scarica", non_valido)
        from tkinter import messagebox
        monkeypatch.setattr(messagebox, "showerror", lambda titolo, testo, **k: errori.append(testo))
        app._apri_aggiornamento()
        assert errori == ["SHA-256 diverso"] and len(avviati) == 1 and chiusa == [True]

        # con articoli in tabella si chiede prima di chiudere
        from processors.models import Articolo
        app.tabella.carica([Articolo(pagina=1, modello="A")])
        monkeypatch.setattr(aggiornamento, "scarica", scarica)
        monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: False)
        app._apri_aggiornamento()
        assert len(avviati) == 1
    finally:
        app.destroy()
    _radice_nuova(radice)


def test_copia_non_installata_apre_la_pagina(radice, tmp_path, monkeypatch):
    import aggiornamento

    app, finestra = _finestra_con_aggiornamento(radice, tmp_path, monkeypatch, _rilascio_di_prova())
    aperte = []

    class Aggiorna(finestra.DialogoAggiornamento):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            assert self.btn_aggiorna.cget("text") == "Apri la pagina di download"
            self.after(100, lambda: self._scegli("aggiorna"))

    monkeypatch.setattr(finestra, "DialogoAggiornamento", Aggiorna)
    monkeypatch.setattr(aggiornamento, "modo_installazione", lambda: None)
    monkeypatch.setattr(finestra.webbrowser, "open", aperte.append)
    try:
        app._apri_aggiornamento()
        assert aperte == ["https://example.com/release"]
    finally:
        app.destroy()
    _radice_nuova(radice)


def test_il_controllo_aspetta_la_fine_dell_estrazione(radice, tmp_path, monkeypatch):
    import aggiornamento

    app, _ = _finestra_con_aggiornamento(radice, tmp_path, monkeypatch, None)
    chieste = []
    monkeypatch.setattr(aggiornamento, "cerca_aggiornamento", lambda: chieste.append(True))
    try:
        app._lavoro = object()  # un'estrazione in corso
        app._controlla_aggiornamento()
        for _ in range(10):
            app.update()
            time.sleep(0.03)
        assert chieste == []
        app._lavoro = None
        app._controlla_aggiornamento()
        limite = time.monotonic() + 5
        while not chieste and time.monotonic() < limite:
            app.update()
            time.sleep(0.03)
        assert chieste == [True]
    finally:
        app._lavoro = None
        app.destroy()
    _radice_nuova(radice)
