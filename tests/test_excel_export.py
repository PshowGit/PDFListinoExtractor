from decimal import Decimal

import pytest
from openpyxl import load_workbook

from exporters.excel_exporter import EsportazioneFallita, esporta_listino, esporta_report_anomalie
from processors.estrazione import estrai_listino
from processors.models import Articolo


def _articoli():
    return [
        Articolo("007-A", "(000123)", "0012345678905", "Prima descrizione", Decimal("1234.5"), 1),
        Articolo("B2", "(B-2)", "", "Seconda", None, 2, note_estrazione={"prezzo": "Prezzo mancante"}),
    ]


def test_struttura_e_formati(tmp_path):
    p = tmp_path / "out.xlsx"
    assert esporta_listino(_articoli(), p) == 2
    wb = load_workbook(p)
    assert wb.sheetnames == ["Listino"]
    ws = wb["Listino"]
    assert [c.value for c in ws[1]] == ["Modello", "Codice", "EAN", "Descrizione", "Prezzo"]
    assert all(c.font.bold for c in ws[1])
    assert ws.max_column == 5 and ws.max_row == 3
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:E3"
    assert ws["D2"].alignment.wrap_text

    modello, codice, ean, _, prezzo = ws[2]
    assert (modello.value, codice.value, ean.value) == ("007-A", "(000123)", "0012345678905")
    assert modello.number_format == codice.number_format == ean.number_format == "@"
    assert isinstance(prezzo.value, float) and prezzo.value == 1234.5
    assert prezzo.number_format == "#,##0.00"
    assert ws["E3"].value is None  # prezzo mancante: cella vuota, non "0"


def test_nessuna_riga_vuota_ne_intestazioni_ripetute(tmp_path, pdf_sintetico, regole):
    arts = estrai_listino(pdf_sintetico, regole).articoli
    p = tmp_path / "s.xlsx"
    esporta_listino(arts, p)
    ws = load_workbook(p)["Listino"]
    righe = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(righe) == len(arts)
    assert all(r[0] or r[1] for r in righe)
    assert not any(r[0] in ("Modello", "Pubblico") for r in righe)


def test_niente_da_esportare(tmp_path):
    with pytest.raises(EsportazioneFallita):
        esporta_listino([], tmp_path / "x.xlsx")


def test_report_anomalie_separato(tmp_path):
    p = tmp_path / "anomalie.xlsx"
    assert esporta_report_anomalie(_articoli(), p) == 1
    ws = load_workbook(p)["Anomalie"]
    assert [c.value for c in ws[2]] == [2, "B2", "(B-2)", "Prezzo mancante"]
